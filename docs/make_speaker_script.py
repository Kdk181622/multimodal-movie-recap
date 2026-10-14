"""Builds docs/Speaker_Script.docx from docs/speaker_notes.json (the same text as the slide speaker notes)."""

import json
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

DOCS = Path(__file__).resolve().parent
PLAYBACK_SLIDE, PLAYBACK_SECONDS = 11, 42  # embedded recap excerpt on slide 11

notes = json.loads((DOCS / "speaker_notes.json").read_text(encoding="utf-8"))
spoken = sum(n["seconds"] for n in notes)
total = spoken + PLAYBACK_SECONDS

doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11.5)
doc.add_heading("Speaker Script: Multimodal Movie Recap", 0)
doc.add_paragraph("AI11 · Module 8 · Assignment 1 · Presenter: Karthikeyan K")
doc.add_paragraph(f"Slides: {len(notes)}  |  Spoken time about {spoken // 60} min {spoken % 60:02d} s, plus the "
                  f"{PLAYBACK_SECONDS} s recap excerpt embedded in slide {PLAYBACK_SLIDE} (total about {total // 60} min "
                  f"{total % 60:02d} s, inside the required 8-20 minutes). The same text is in each slide's speaker notes.")

doc.add_heading("Recording steps", 2)
for tip in [
    "Open docs/Movie_Recap_Presentation.pptx in PowerPoint and choose Slide Show > Record > From Beginning.",
    "Turn the camera on if you want your face in the corner; the microphone records your narration.",
    f"On slide {PLAYBACK_SLIDE}, click the video and let the {PLAYBACK_SECONDS} s excerpt play: PowerPoint records it with its sound.",
    "Speak at a relaxed pace (about 130 words per minute); pause briefly when moving to the next slide.",
    "When finished: File > Export > Create a Video (Full HD, use recorded timings and narrations) to get an MP4.",
    "Check the exported MP4 is between 8 and 20 minutes before uploading.",
]:
    doc.add_paragraph(tip, style="List Bullet")

clock = 0
for i, note in enumerate(notes, 1):
    doc.add_heading(f"Slide {i}: {note['title']}", 2)
    end = clock + note["seconds"]
    meta = doc.add_paragraph().add_run(
        f"{clock // 60}:{clock % 60:02d} - {end // 60}:{end % 60:02d}  (~{note['seconds']} s, {len(note['script'].split())} words)")
    meta.italic = True
    meta.font.color.rgb = RGBColor(0x8A, 0x89, 0x84)
    doc.add_paragraph(note["script"])
    clock = end + (PLAYBACK_SECONDS if i == PLAYBACK_SLIDE else 0)

out = DOCS / "Speaker_Script.docx"
doc.save(out)
print(f"wrote {out.name}: {spoken} s spoken + {PLAYBACK_SECONDS} s playback = {total} s")
